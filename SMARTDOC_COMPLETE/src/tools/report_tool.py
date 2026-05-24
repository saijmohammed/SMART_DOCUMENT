"""
src/tools/report_tool.py
Génération de rapports PDF et Word professionnels pour SmartDoc AI.
"""
from __future__ import annotations

import html
import re
from datetime import datetime
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    HRFlowable, PageBreak, Paragraph, SimpleDocTemplate,
    Spacer, Table, TableStyle,
)

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor, Cm, Inches
    DOCX_AVAILABLE = True
except Exception:
    DOCX_AVAILABLE = False

ROOT        = Path(__file__).resolve().parents[2]
REPORTS_DIR = ROOT / "reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

# ── Palette couleurs ──────────────────────────────────────────────────────────
C_DARK    = colors.HexColor("#0f172a")
C_BLUE    = colors.HexColor("#1d4ed8")
C_LBLUE   = colors.HexColor("#dbeafe")
C_ACCENT  = colors.HexColor("#3b82f6")
C_GREY    = colors.HexColor("#64748b")
C_LGREY   = colors.HexColor("#f1f5f9")
C_WHITE   = colors.white
C_BORDER  = colors.HexColor("#cbd5e1")
C_GREEN   = colors.HexColor("#059669")
C_LGREEN  = colors.HexColor("#d1fae5")
C_VIOLET  = colors.HexColor("#7c3aed")
C_LVIOLET = colors.HexColor("#ede9fe")


# =============================================================================
# Helpers texte
# =============================================================================

def _now_slug() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _clean(text) -> str:
    if not text:
        return ""
    text = str(text)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    text = html.unescape(text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*",     r"\1", text)
    return re.sub(r"\n{3,}", "\n\n", text.replace("\r", "\n")).strip()


def _esc(text: str) -> str:
    return html.escape(_clean(text)).replace("\n", "<br/>")


def _doc_name(result: dict) -> str:
    return (result.get("document", {}).get("file_name")
            or result.get("file_name") or "document")


def _summary_text(result: dict) -> str:
    return _clean(result.get("summary", {}).get("summary", ""))


def _key_points(result: dict) -> list[str]:
    raw = result.get("summary", {}).get("key_points", [])
    return [_clean(k) for k in raw if _clean(k)]


def _explanation_text(result: dict) -> str:
    return _clean(result.get("explanation", {}).get("simple_explanation", ""))


def _terms(result: dict) -> list[str]:
    return result.get("explanation", {}).get("important_terms", [])


def _questions(result: dict) -> list[str]:
    return [_clean(q) for q in result.get("explanation", {}).get("study_questions", []) if _clean(q)]


def _translation_text(result: dict) -> str:
    return _clean(result.get("translation", {}).get("translated_text", ""))


_LANG_NAMES = {
    "fr":"Français","en":"English","ar":"Arabe","es":"Español",
    "de":"Deutsch","it":"Italiano","pt":"Português","zh":"中文",
    "ja":"日本語","ru":"Русский","tr":"Türkçe","nl":"Nederlands",
}


def _translation_language(result: dict) -> str:
    code = result.get("translation", {}).get("target_language", "")
    return _LANG_NAMES.get(code, code or "—")


def _doc_meta(result: dict) -> dict:
    doc = result.get("document", {})
    cls = result.get("classification", {})
    return {
        "title"   : result.get("summary", {}).get("title", "") or _doc_name(result),
        "pages"   : doc.get("page_count", "—"),
        "words"   : doc.get("word_count", 0),
        "category": cls.get("category", "—"),
        "file"    : _doc_name(result),
        "date"    : datetime.now().strftime("%d/%m/%Y à %H:%M"),
    }


# =============================================================================
# PDF — styles
# =============================================================================

def _pdf_styles():
    styles = getSampleStyleSheet()

    defs = [
        ("Cover_Title", {
            "fontName":"Helvetica-Bold","fontSize":26,"leading":32,
            "textColor":C_WHITE,"alignment":TA_CENTER,"spaceAfter":6,
        }),
        ("Cover_Sub", {
            "fontName":"Helvetica","fontSize":13,"leading":18,
            "textColor":colors.HexColor("#bfdbfe"),"alignment":TA_CENTER,"spaceAfter":4,
        }),
        ("Cover_Meta", {
            "fontName":"Helvetica","fontSize":10,"leading":14,
            "textColor":colors.HexColor("#94a3b8"),"alignment":TA_CENTER,
        }),
        ("Section_H", {
            "fontName":"Helvetica-Bold","fontSize":13,"leading":17,
            "textColor":C_WHITE,"alignment":TA_LEFT,
            "spaceBefore":0,"spaceAfter":0,
            "leftIndent":0,
        }),
        ("Body_J", {
            "fontName":"Helvetica","fontSize":10.5,"leading":17,
            "textColor":C_DARK,"alignment":TA_JUSTIFY,
            "spaceBefore":4,"spaceAfter":6,
        }),
        ("Bullet_Item", {
            "fontName":"Helvetica","fontSize":10.5,"leading":16,
            "textColor":C_DARK,"leftIndent":16,"bulletIndent":4,
            "spaceBefore":2,"spaceAfter":2,
        }),
        ("Note_Style", {
            "fontName":"Helvetica-Oblique","fontSize":10,"leading":15,
            "textColor":colors.HexColor("#374151"),
            "backColor":colors.HexColor("#fefce8"),
            "borderColor":colors.HexColor("#fbbf24"),
            "borderWidth":1,"borderPadding":10,
            "spaceBefore":6,"spaceAfter":10,
        }),
        ("Meta_Label", {
            "fontName":"Helvetica-Bold","fontSize":9,
            "textColor":C_BLUE,"alignment":TA_LEFT,
        }),
        ("Meta_Value", {
            "fontName":"Helvetica","fontSize":9,
            "textColor":C_DARK,"alignment":TA_LEFT,
        }),
        ("Footer_S", {
            "fontName":"Helvetica","fontSize":8,
            "textColor":C_GREY,"alignment":TA_CENTER,
        }),
        ("Tag_S", {
            "fontName":"Helvetica-Bold","fontSize":9,
            "textColor":C_VIOLET,"alignment":TA_CENTER,
        }),
    ]
    for name, cfg in defs:
        styles.add(ParagraphStyle(name=name, parent=styles["Normal"], **cfg))
    return styles


def _pdf_doc(path: Path, title: str):
    return SimpleDocTemplate(
        str(path), pagesize=A4,
        rightMargin=1.9*cm, leftMargin=1.9*cm,
        topMargin=1.6*cm,   bottomMargin=1.8*cm,
        title=title,
    )


def _section_block(story, styles, heading: str, color=C_BLUE):
    """Bloc section avec fond coloré."""
    tbl = Table([[Paragraph(f"▌  {heading}", styles["Section_H"])]],
                colWidths=[17.2*cm])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), color),
        ("LEFTPADDING",  (0,0), (-1,-1), 12),
        ("RIGHTPADDING", (0,0), (-1,-1), 12),
        ("TOPPADDING",   (0,0), (-1,-1), 8),
        ("BOTTOMPADDING",(0,0), (-1,-1), 8),
        ("ROUNDEDCORNERS", [4]),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 8))


def _meta_table(story, styles, meta: dict):
    rows = [
        [Paragraph("Titre du document", styles["Meta_Label"]),
         Paragraph(_esc(meta["title"]), styles["Meta_Value"])],
        [Paragraph("Fichier source",    styles["Meta_Label"]),
         Paragraph(_esc(meta["file"]),  styles["Meta_Value"])],
        [Paragraph("Catégorie",         styles["Meta_Label"]),
         Paragraph(_esc(meta["category"]), styles["Meta_Value"])],
        [Paragraph("Volume",            styles["Meta_Label"]),
         Paragraph(f"{meta['pages']} pages  ·  {meta['words']:,} mots", styles["Meta_Value"])],
        [Paragraph("Date de génération",styles["Meta_Label"]),
         Paragraph(meta["date"],        styles["Meta_Value"])],
        [Paragraph("Plateforme",        styles["Meta_Label"]),
         Paragraph("SmartDoc AI v4.0", styles["Meta_Value"])],
    ]
    tbl = Table(rows, colWidths=[4.5*cm, 12.7*cm])
    tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0,0), (0,-1), C_LBLUE),
        ("BACKGROUND",    (1,0), (1,-1), C_LGREY),
        ("GRID",          (0,0), (-1,-1), 0.4, C_BORDER),
        ("TOPPADDING",    (0,0), (-1,-1), 7),
        ("BOTTOMPADDING", (0,0), (-1,-1), 7),
        ("LEFTPADDING",   (0,0), (-1,-1), 10),
        ("RIGHTPADDING",  (0,0), (-1,-1), 10),
        ("VALIGN",        (0,0), (-1,-1), "TOP"),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 16))


def _pdf_cover(story, styles, report_title: str, meta: dict):
    """Page de couverture avec fond sombre."""
    cover_tbl = Table(
        [[Paragraph("SmartDoc AI", styles["Cover_Sub"]),
          Paragraph(report_title,  styles["Cover_Title"]),
          Paragraph(f"Document : {meta['title']}", styles["Cover_Sub"]),
          Paragraph(f"Généré le {meta['date']}",   styles["Cover_Meta"])]],
        colWidths=[17.2*cm],
    )
    cover_tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0,0), (-1,-1), C_DARK),
        ("TOPPADDING",    (0,0), (-1,-1), 30),
        ("BOTTOMPADDING", (0,0), (-1,-1), 30),
        ("LEFTPADDING",   (0,0), (-1,-1), 20),
        ("RIGHTPADDING",  (0,0), (-1,-1), 20),
        ("ROUNDEDCORNERS", [8]),
    ]))
    story.append(cover_tbl)
    story.append(Spacer(1, 20))


def _body_paragraphs(story, styles, text: str):
    text = _clean(text) or "Non disponible."
    for para in text.split("\n\n"):
        para = para.strip()
        if para:
            story.append(Paragraph(html.escape(para).replace("\n","<br/>"),
                                   styles["Body_J"]))
    story.append(Spacer(1, 6))


def _bullet_list(story, styles, items: list[str], color=C_BLUE):
    for item in items:
        item = _clean(item)
        if item:
            row = Table(
                [[Paragraph("●", ParagraphStyle("_dot", parent=styles["Normal"],
                    fontName="Helvetica-Bold", fontSize=11, textColor=color)),
                  Paragraph(html.escape(item), styles["Bullet_Item"])]],
                colWidths=[0.5*cm, 16.2*cm],
            )
            row.setStyle(TableStyle([
                ("VALIGN",       (0,0),(-1,-1),"TOP"),
                ("TOPPADDING",   (0,0),(-1,-1),2),
                ("BOTTOMPADDING",(0,0),(-1,-1),2),
                ("LEFTPADDING",  (0,0),(-1,-1),0),
            ]))
            story.append(row)
    story.append(Spacer(1, 6))


def _numbered_list(story, styles, items: list[str]):
    for i, item in enumerate(items, 1):
        item = _clean(item)
        if item:
            row = Table(
                [[Paragraph(f"Q{i}.", ParagraphStyle("_num", parent=styles["Normal"],
                    fontName="Helvetica-Bold", fontSize=10, textColor=C_VIOLET)),
                  Paragraph(html.escape(item), styles["Bullet_Item"])]],
                colWidths=[0.9*cm, 15.8*cm],
            )
            row.setStyle(TableStyle([
                ("VALIGN",       (0,0),(-1,-1),"TOP"),
                ("TOPPADDING",   (0,0),(-1,-1),3),
                ("BOTTOMPADDING",(0,0),(-1,-1),3),
            ]))
            story.append(row)
    story.append(Spacer(1, 6))


def _tags_row(story, styles, terms: list[str]):
    if not terms:
        return
    joined = "   ·   ".join(f"[ {t} ]" for t in terms[:12])
    story.append(Paragraph(joined, styles["Tag_S"]))
    story.append(Spacer(1, 8))


def _divider(story):
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=0.5,
                             color=C_BORDER, spaceAfter=10))


# =============================================================================
# PDF — rapport résumé
# =============================================================================

def generate_summary_pdf_report(result: dict, remarks: str = "") -> dict:
    path   = REPORTS_DIR / f"smartdoc_summary_{_now_slug()}.pdf"
    styles = _pdf_styles()
    meta   = _doc_meta(result)
    story  = []

    _pdf_cover(story, styles, "Résumé détaillé du document", meta)

    # Métadonnées
    _section_block(story, styles, "Informations sur le document", C_DARK)
    _meta_table(story, styles, meta)

    # Résumé
    _section_block(story, styles, "Résumé détaillé", C_BLUE)
    _body_paragraphs(story, styles, _summary_text(result))
    _divider(story)

    # Points clés
    kps = _key_points(result)
    if kps:
        _section_block(story, styles, "Points clés", C_GREEN)
        _bullet_list(story, styles, kps, C_GREEN)
        _divider(story)

    # Remarques
    if remarks:
        _section_block(story, styles, "Remarques personnelles",
                       colors.HexColor("#d97706"))
        story.append(Paragraph(_esc(remarks), styles["Note_Style"]))

    _pdf_doc(path, "Résumé détaillé").build(story)
    return {"status": "success", "report_path": str(path)}


# =============================================================================
# PDF — rapport traduction
# =============================================================================

def generate_translation_pdf_report(result: dict) -> dict:
    path   = REPORTS_DIR / f"smartdoc_translation_{_now_slug()}.pdf"
    styles = _pdf_styles()
    meta   = _doc_meta(result)
    lang   = _translation_language(result)
    story  = []

    _pdf_cover(story, styles, f"Traduction — {lang}", meta)
    _section_block(story, styles, "Informations", C_DARK)
    _meta_table(story, styles, meta)

    _section_block(story, styles, "Texte source (résumé)", C_BLUE)
    _body_paragraphs(story, styles, _summary_text(result))
    _divider(story)

    story.append(PageBreak())
    _section_block(story, styles, f"Traduction en {lang}", C_VIOLET)
    _body_paragraphs(story, styles, _translation_text(result))

    _pdf_doc(path, f"Traduction — {lang}").build(story)
    return {"status": "success", "report_path": str(path)}


# =============================================================================
# PDF — rapport général
# =============================================================================

def generate_general_pdf_report(
    result: dict, remarks: str = "", validator_note: str = ""
) -> dict:
    path   = REPORTS_DIR / f"smartdoc_general_{_now_slug()}.pdf"
    styles = _pdf_styles()
    meta   = _doc_meta(result)
    story  = []

    _pdf_cover(story, styles, "Rapport d'Analyse Documentaire", meta)

    # Métadonnées
    _section_block(story, styles, "Informations sur le document", C_DARK)
    _meta_table(story, styles, meta)

    # Résumé
    _section_block(story, styles, "Résumé essentiel", C_BLUE)
    _body_paragraphs(story, styles, _summary_text(result))

    kps = _key_points(result)
    if kps:
        story.append(Paragraph("<b>Points clés :</b>", styles["Body_J"]))
        _bullet_list(story, styles, kps, C_BLUE)
    _divider(story)

    # Explication
    expl = _explanation_text(result)
    if expl:
        story.append(PageBreak())
        _section_block(story, styles, "Explication simplifiée", C_GREEN)
        _body_paragraphs(story, styles, expl)

        tms = _terms(result)
        if tms:
            story.append(Paragraph("<b>Termes importants :</b>", styles["Body_J"]))
            _tags_row(story, styles, tms)

        qs = _questions(result)
        if qs:
            story.append(Paragraph("<b>Questions de compréhension :</b>", styles["Body_J"]))
            _numbered_list(story, styles, qs)
        _divider(story)

    # Traduction
    transl = _translation_text(result)
    if transl:
        story.append(PageBreak())
        lang = _translation_language(result)
        _section_block(story, styles, f"Traduction — {lang}", C_VIOLET)
        _body_paragraphs(story, styles, transl)
        _divider(story)

    # Remarques
    if remarks:
        _section_block(story, styles, "Remarques du lecteur",
                       colors.HexColor("#d97706"))
        story.append(Paragraph(_esc(remarks), styles["Note_Style"]))
        _divider(story)

    # Validation HITL
    note = validator_note or "Rapport validé par l'utilisateur."
    _section_block(story, styles, "Validation humaine (HITL)",
                   colors.HexColor("#065f46"))
    story.append(Paragraph(_esc(note), styles["Note_Style"]))

    _pdf_doc(path, "Rapport d'Analyse Documentaire").build(story)
    return {"status": "success", "report_path": str(path)}


# =============================================================================
# Word — helpers
# =============================================================================

def _shade_cell(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color.lstrip("#"))
    tcPr.append(shd)


def _docx_cover(doc, report_title: str, meta: dict):
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SmartDoc AI")
    r.font.size  = Pt(11); r.font.color.rgb = RGBColor(100,116,139)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(report_title)
    r.bold = True; r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(15,23,42)

    doc.add_paragraph("")
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    cell = tbl.cell(0,0)
    _shade_cell(cell, "#dbeafe")
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(
        f"{meta['title']}\n"
        f"Catégorie : {meta['category']}  ·  {meta['pages']} pages  ·  {meta['words']:,} mots\n"
        f"Généré le {meta['date']}"
    )
    cr.font.size = Pt(10)
    cr.font.color.rgb = RGBColor(30,64,175)
    doc.add_paragraph("")


def _docx_section(doc, heading: str, hex_bg="#1d4ed8", hex_fg="ffffff"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    cell = tbl.cell(0,0)
    _shade_cell(cell, hex_bg)
    cp = cell.paragraphs[0]
    r  = cp.add_run(f"  {heading}")
    r.bold = True; r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(
        int(hex_fg[0:2],16), int(hex_fg[2:4],16), int(hex_fg[4:6],16)
    )
    doc.add_paragraph("")


def _docx_meta_table(doc, meta: dict):
    rows = [
        ("Titre du document", meta["title"]),
        ("Fichier source",    meta["file"]),
        ("Catégorie",         meta["category"]),
        ("Volume",            f"{meta['pages']} pages  ·  {meta['words']:,} mots"),
        ("Date de génération",meta["date"]),
        ("Plateforme",        "SmartDoc AI v4.0"),
    ]
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.style = "Table Grid"
    for i, (label, value) in enumerate(rows):
        cell_l = tbl.cell(i, 0)
        cell_r = tbl.cell(i, 1)
        _shade_cell(cell_l, "dbeafe")
        _shade_cell(cell_r, "f8fafc")
        rl = cell_l.paragraphs[0].add_run(label)
        rl.bold = True; rl.font.size = Pt(9.5)
        rl.font.color.rgb = RGBColor(29,78,216)
        rr = cell_r.paragraphs[0].add_run(value)
        rr.font.size = Pt(9.5)
    tbl.columns[0].width = Cm(4.5)
    tbl.columns[1].width = Cm(12)
    doc.add_paragraph("")


def _docx_body(doc, text: str):
    text = _clean(text) or "Non disponible."
    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after  = Pt(6)
        r = p.add_run(para.replace("\n", " "))
        r.font.size = Pt(11)


def _docx_bullets(doc, items: list[str], hex_color="1d4ed8"):
    for item in items:
        item = _clean(item)
        if not item:
            continue
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.space_after  = Pt(3)
        r = p.add_run(item)
        r.font.size = Pt(11)
    doc.add_paragraph("")


def _docx_numbered(doc, items: list[str]):
    for i, item in enumerate(items, 1):
        item = _clean(item)
        if not item:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.5)
        p.paragraph_format.space_after  = Pt(4)
        rn = p.add_run(f"Q{i}.  ")
        rn.bold = True; rn.font.size = Pt(11)
        rn.font.color.rgb = RGBColor(124,58,237)
        r = p.add_run(item)
        r.font.size = Pt(11)
    doc.add_paragraph("")


def _docx_tags(doc, terms: list[str]):
    if not terms:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    for i, t in enumerate(terms[:12]):
        r = p.add_run(f"[ {t} ]")
        r.bold = True; r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(124,58,237)
        if i < len(terms)-1:
            p.add_run("   ").font.size = Pt(9)
    doc.add_paragraph("")


def _docx_note(doc, text: str, hex_bg="fefce8"):
    text = _clean(text) or "Non disponible."
    tbl  = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    cell = tbl.cell(0,0)
    _shade_cell(cell, hex_bg)
    cp = cell.paragraphs[0]
    cp.paragraph_format.left_indent  = Cm(0.3)
    cp.paragraph_format.space_after  = Pt(4)
    r  = cp.add_run(text)
    r.font.size   = Pt(10.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(55,65,81)
    doc.add_paragraph("")


def _docx_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pb  = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"),  "single")
    bot.set(qn("w:sz"),   "4")
    bot.set(qn("w:space"),"1")
    bot.set(qn("w:color"),"cbd5e1")
    pb.append(bot)
    pPr.append(pb)
    doc.add_paragraph("")


# =============================================================================
# Word — rapports publics
# =============================================================================

def generate_summary_docx_report(result: dict, remarks: str = "") -> dict:
    if not DOCX_AVAILABLE:
        return {"status":"error","error":"python-docx non installé."}
    path = REPORTS_DIR / f"smartdoc_summary_{_now_slug()}.docx"
    meta = _doc_meta(result)
    doc  = Document()

    _docx_cover(doc, "Résumé détaillé du document", meta)
    _docx_section(doc, "Informations sur le document", "0f172a")
    _docx_meta_table(doc, meta)
    _docx_divider(doc)

    _docx_section(doc, "Résumé détaillé", "1d4ed8")
    _docx_body(doc, _summary_text(result))
    _docx_divider(doc)

    kps = _key_points(result)
    if kps:
        _docx_section(doc, "Points clés", "059669", "ffffff")
        _docx_bullets(doc, kps)
        _docx_divider(doc)

    if remarks:
        _docx_section(doc, "Remarques personnelles", "d97706")
        _docx_note(doc, remarks)

    doc.save(str(path))
    return {"status":"success","report_path":str(path)}


def generate_translation_docx_report(result: dict) -> dict:
    if not DOCX_AVAILABLE:
        return {"status":"error","error":"python-docx non installé."}
    path = REPORTS_DIR / f"smartdoc_translation_{_now_slug()}.docx"
    lang = _translation_language(result)
    meta = _doc_meta(result)
    doc  = Document()

    _docx_cover(doc, f"Traduction — {lang}", meta)
    _docx_section(doc, "Informations", "0f172a")
    _docx_meta_table(doc, meta)
    _docx_divider(doc)

    _docx_section(doc, "Texte source (résumé)", "1d4ed8")
    _docx_body(doc, _summary_text(result))
    _docx_divider(doc)

    _docx_section(doc, f"Traduction en {lang}", "7c3aed", "ffffff")
    _docx_body(doc, _translation_text(result))

    doc.save(str(path))
    return {"status":"success","report_path":str(path)}


def generate_general_docx_report(
    result: dict, remarks: str = "", validator_note: str = ""
) -> dict:
    if not DOCX_AVAILABLE:
        return {"status":"error","error":"python-docx non installé."}
    path = REPORTS_DIR / f"smartdoc_general_{_now_slug()}.docx"
    meta = _doc_meta(result)
    doc  = Document()

    _docx_cover(doc, "Rapport d'Analyse Documentaire", meta)

    _docx_section(doc, "Informations sur le document", "0f172a")
    _docx_meta_table(doc, meta)
    _docx_divider(doc)

    _docx_section(doc, "Résumé essentiel", "1d4ed8")
    _docx_body(doc, _summary_text(result))
    kps = _key_points(result)
    if kps:
        p = doc.add_paragraph()
        r = p.add_run("Points clés :")
        r.bold = True; r.font.size = Pt(11)
        _docx_bullets(doc, kps)
    _docx_divider(doc)

    expl = _explanation_text(result)
    if expl:
        _docx_section(doc, "Explication simplifiée", "059669")
        _docx_body(doc, expl)
        tms = _terms(result)
        if tms:
            p = doc.add_paragraph()
            r = p.add_run("Termes importants :")
            r.bold = True; r.font.size = Pt(11)
            _docx_tags(doc, tms)
        qs = _questions(result)
        if qs:
            p = doc.add_paragraph()
            r = p.add_run("Questions de compréhension :")
            r.bold = True; r.font.size = Pt(11)
            _docx_numbered(doc, qs)
        _docx_divider(doc)

    transl = _translation_text(result)
    if transl:
        lang = _translation_language(result)
        _docx_section(doc, f"Traduction — {lang}", "7c3aed", "ffffff")
        _docx_body(doc, transl)
        _docx_divider(doc)

    if remarks:
        _docx_section(doc, "Remarques du lecteur", "d97706")
        _docx_note(doc, remarks)
        _docx_divider(doc)

    note = validator_note or "Rapport validé par l'utilisateur."
    _docx_section(doc, "Validation humaine (HITL)", "065f46", "ffffff")
    _docx_note(doc, note, "d1fae5")

    doc.save(str(path))
    return {"status":"success","report_path":str(path)}


# =============================================================================
# Rétro-compatibilité
# =============================================================================

def generate_pdf_report(result: dict, validator_note: str = "", **kwargs) -> dict:
    return generate_general_pdf_report(result, validator_note=validator_note)


def generate_report(result: dict, validator_note: str = "", **kwargs) -> dict:
    return generate_general_pdf_report(result, validator_note=validator_note)
